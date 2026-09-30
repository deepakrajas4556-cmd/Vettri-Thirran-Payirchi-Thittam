from io import BytesIO
import re
from typing import Optional

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from fpdf import FPDF

from document_utils.sanitize import sanitize_text


def _split_lines(text: str) -> list[str]:
    return [line.strip() for line in sanitize_text(text).splitlines()]


def _safe_filename(document_type: str, extension: str) -> str:
    base = re.sub(r"[^A-Za-z0-9_-]+", "_", document_type.strip()).strip("_")
    base = base or "legal_document"
    return f"{base}.{extension}"


def format_txt(text: str) -> bytes:
    return sanitize_text(text).encode("utf-8")


def format_docx(
    text: str,
    document_type: str,
    logo_bytes: Optional[bytes] = None,
    terms: Optional[str] = None,
) -> tuple[bytes, str]:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)

    styles = doc.styles
    styles["Normal"].font.name = "Times New Roman"
    styles["Normal"].font.size = Pt(11)

    if logo_bytes:
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run()
        run.add_picture(BytesIO(logo_bytes), width=Inches(1.2))

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(document_type.strip().upper())
    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(16)

    for line in _split_lines(text):
        if not line:
            doc.add_paragraph()
            continue

        clean = line.strip()
        is_heading = (
            len(clean) < 100
            and (
                clean.isupper()
                or bool(re.match(r"^(SECTION|ARTICLE|[0-9]+[.)])\s", clean, re.I))
            )
        )

        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(6)
        if is_heading:
            run = paragraph.add_run(clean)
            run.bold = True
        else:
            paragraph.add_run(clean)

    if terms:
        items = [x.strip() for x in terms.split(";") if x.strip()]
        if items:
            doc.add_paragraph()
            heading = doc.add_paragraph()
            heading.add_run("Key Terms").bold = True
            table = doc.add_table(rows=1, cols=2)
            table.style = "Table Grid"
            table.rows[0].cells[0].text = "No."
            table.rows[0].cells[1].text = "Term"
            for index, item in enumerate(items, start=1):
                cells = table.add_row().cells
                cells[0].text = str(index)
                cells[1].text = item

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("LegalEase - AI-assisted draft. Review before execution.")

    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue(), _safe_filename(document_type, "docx")


class _LegalEasePDF(FPDF):
    def __init__(self, logo_bytes: Optional[bytes] = None):
        super().__init__()
        self.logo_bytes = logo_bytes

    def header(self):
        if self.logo_bytes:
            try:
                self.image(BytesIO(self.logo_bytes), x=95, y=8, w=20)
                self.set_y(31)
            except Exception:
                self.set_y(12)
        else:
            self.set_y(12)
        self.set_font("Helvetica", "B", 8)
        self.cell(0, 6, "LEGALEASE", align="C")
        self.ln(8)

    def footer(self):
        self.set_y(-15)
        self.set_font("Helvetica", "", 8)
        self.cell(0, 8, "LegalEase - AI-assisted draft. Review before execution.", align="C")


def format_pdf(
    text: str,
    document_type: str,
    logo_bytes: Optional[bytes] = None,
) -> tuple[bytes, str]:
    pdf = _LegalEasePDF(logo_bytes=logo_bytes)
    pdf.set_auto_page_break(auto=True, margin=20)
    pdf.add_page()

    pdf.set_font("Helvetica", "B", 16)
    pdf.multi_cell(0, 9, document_type.strip().upper(), align="C")
    pdf.ln(4)

    for line in _split_lines(text):
        if not line:
            pdf.ln(4)
            continue

        is_heading = (
            len(line) < 100
            and (
                line.isupper()
                or bool(re.match(r"^(SECTION|ARTICLE|[0-9]+[.)])\s", line, re.I))
            )
        )
        pdf.set_font("Helvetica", "B" if is_heading else "", 11)
        pdf.multi_cell(0, 6, line)
        pdf.ln(1)

    output = pdf.output()
    return bytes(output), _safe_filename(document_type, "pdf")
